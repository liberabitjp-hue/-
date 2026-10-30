"""拡大機（ポスタープリンター）用の原稿を作成する。

A4縦・1ページ1文字・上下左右中央・黒100%で、次の2ファイルを出力する。
  - PDF : 印刷用。フォント埋め込み・文字の正確な中央配置
  - Word: 編集用。Windows に BIZ UDゴシック が入っていれば同じ見た目になる
使い方: python3 make_poster.py  （BIZUDGothic-Bold.ttf を FONT_FILE に置くこと）
"""
from pathlib import Path

from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

TEXT = "お米マーク総選挙"
FONT = "BIZ UDGothic"  # 日本語表示名: BIZ UDゴシック（Bold で使用）
FONT_FILE = Path.home() / ".fonts" / "BIZUDGothic-Bold.ttf"
FONT_SIZE_PT = 500  # 全角1文字 ≒ 176mm 角
MARGIN_MM = 15
PAGE_W_MM, PAGE_H_MM = 210, 297
OUT_DIR = Path(__file__).parent
DOCX_PATH = OUT_DIR / f"{TEXT}.docx"
PDF_PATH = OUT_DIR / f"{TEXT}.pdf"


def set_font(run):
    run.font.name = FONT
    run.font.size = Pt(FONT_SIZE_PT)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def build_docx():
    doc = Document()

    section = doc.sections[0]
    section.page_width, section.page_height = Mm(PAGE_W_MM), Mm(PAGE_H_MM)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Mm(MARGIN_MM))
    section.header_distance = section.footer_distance = Mm(0)

    # 上下中央: 行の高さ＝文字サイズとし、残りの半分を段落前の余白にする
    # （セクションの上下中央揃えは LibreOffice が反映しないため使わない）
    # 固定行間では文字が行の上側に寄るため、実測（LibreOffice 変換）で 14mm 下げる
    line_mm = FONT_SIZE_PT * 25.4 / 72
    space_before_mm = (PAGE_H_MM - 2 * MARGIN_MM - line_mm) / 2 + 14

    # 既定の空段落を削除
    body = doc.element.body
    for p in list(body.iterchildren(qn("w:p"))):
        body.remove(p)

    for i, char in enumerate(TEXT):
        para = doc.add_paragraph()
        fmt = para.paragraph_format
        fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.space_before = Mm(space_before_mm)
        fmt.space_after = Pt(0)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(FONT_SIZE_PT)
        fmt.page_break_before = i > 0
        # 文字グリッドへの吸着を無効化（日本語版 Word で位置がずれないように）
        snap = OxmlElement("w:snapToGrid")
        snap.set(qn("w:val"), "0")
        para._p.get_or_add_pPr().append(snap)
        set_font(para.add_run(char))

    doc.save(DOCX_PATH)
    print(f"saved: {DOCX_PATH} ({len(TEXT)} pages)")


def build_pdf():
    font = TTFont("BIZUDGothic-Bold", str(FONT_FILE))
    pdfmetrics.registerFont(font)
    # 全角の仮想ボディ（ascent〜descent）の中心をページ中心に合わせる
    em_center = (font.face.ascent + font.face.descent) / 2 / 1000 * FONT_SIZE_PT
    page_w, page_h = PAGE_W_MM * mm, PAGE_H_MM * mm

    c = canvas.Canvas(
        str(PDF_PATH), pagesize=(page_w, page_h), initialFontName=font.fontName
    )
    c.setTitle(TEXT)
    for char in TEXT:
        c.setFillColorRGB(0, 0, 0)
        c.setFont(font.fontName, FONT_SIZE_PT)
        c.drawCentredString(page_w / 2, page_h / 2 - em_center, char)
        c.showPage()
    c.save()
    print(f"saved: {PDF_PATH} ({len(TEXT)} pages)")


def main():
    build_docx()
    build_pdf()


if __name__ == "__main__":
    main()
